from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from typing import Any

from openpyxl import load_workbook

from .base import BankParser, RawTxn, StatementData


class GolomtParser(BankParser):
    code = "GOLOMT"
    name = "Golomt Bank"

    @classmethod
    def detect(cls, workbook: Any) -> bool:
        try:
            ws = workbook.active
            headers = [str(cell.value).strip().lower() for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        except Exception:
            return False
        if not headers:
            return False
        return "date" in headers and "description" in headers and "amount" in headers

    def parse(self, workbook: Any) -> StatementData:
        ws = workbook.active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        account_no = ""
        for row in rows[:20]:
            if row and any(cell is not None for cell in row):
                if len(row) > 3 and isinstance(row[0], str):
                    account_no = row[0].strip() or account_no
                    break
        txns: list[RawTxn] = []
        for row in rows:
            if not row or not any(cell is not None for cell in row):
                continue
            if len(row) < 8:
                continue
            date_value, description, amount_value, *rest = row
            if date_value is None or description is None or amount_value is None:
                continue
            try:
                dt = datetime.strptime(str(date_value), "%Y-%m-%d")
            except ValueError:
                continue
            try:
                amount = Decimal(str(amount_value).replace(",", "").strip())
            except Exception:
                continue
            direction = "in" if amount >= 0 else "out"
            txns.append(
                RawTxn(
                    posted_at=dt,
                    txn_at=dt,
                    amount=abs(amount),
                    direction=direction,
                    description=str(description).strip(),
                    merchant=str(description).strip(),
                    txn_type="other",
                    raw={"source_row": list(row)},
                )
            )
        return StatementData(
            bank_code=self.code,
            bank_name=self.name,
            account_no=account_no or "UNKNOWN",
            transactions=txns,
        )
