from __future__ import annotations

from collections.abc import Iterator
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from collections import Counter

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker

from core.enrich import enrich, link_fees, normalize_account
from core.importer import import_file
from core.models import Base, Category, Rule, Transaction
from core.parsers.golomt import GolomtParser
from core.parsers.registry import detect_parser


def _statement_bytes(report_date: str = "2026-01-14") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "OperativeAccountStatementRaw"
    sheet.append(
        [
            "Хуулганы огноо",
            report_date,
            "ДАНСНЫ ХУУЛГА",
            "Харилцагчийн нэр",
            "ТЕСТ ",
            "Гүйлгээний огноо",
            "2026-01-01 - 2026-01-14",
            "Дансны дугаар",
            "1800000001 [MNT]",
            "IBAN",
            "MN00 0015 0018 0000 0001",
        ]
    )
    sheet.append(
        [
            "Гүйлгээний огноо",
            "Эхний үлдэгдэл",
            "Зарлага",
            "Орлого",
            "Эцсийн үлдэгдэл",
            "Гүйлгээний утга",
            "Харьцсан данс",
        ]
    )
    sheet.append(
        [
            "2026-01-02T10:00:00",
            100.00,
            10.00,
            "",
            90.00,
            "SocialPay гүйлгээ,ЦЭЦЭГ БОЛД,ХАС",
            "MN44 0032 0050 0600 0001",
        ]
    )
    sheet.append(
        [
            "2026-01-02T10:00:00",
            90.00,
            1.00,
            "",
            89.00,
            "Charges for PORD Customer Payment :000658547340",
            "MN440032005006000001",
        ]
    )
    sheet.append(
        [
            "2026-01-02T10:03:00",
            89.00,
            2.00,
            "",
            87.00,
            "400000******1111:02-01-2026 10:02:00:POS:TEST STORE",
            "",
        ]
    )
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _full_sample_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "OperativeAccountStatementRaw"
    sheet.append(
        [
            "Хуулганы огноо",
            "2026-10-14",
            "ДАНСНЫ ХУУЛГА",
            "Харилцагчийн нэр",
            "ТЕСТ ",
            "Гүйлгээний огноо",
            "2026-09-14 - 2026-10-14",
            "Дансны дугаар",
            "1800000001 [MNT]",
            "IBAN",
            "MN00 0015 0018 0000 0001",
        ]
    )
    sheet.append(
        [
            "Гүйлгээний огноо",
            "Эхний үлдэгдэл",
            "Зарлага",
            "Орлого",
            "Эцсийн үлдэгдэл",
            "Гүйлгээний утга",
            "Харьцсан данс",
        ]
    )
    rows = [
        ["2026-09-14T13:37:46", 150000.00, "", 20000, 170000.00, "O-ТЕСТ ХЭРЭГЛЭГЧ", "5000000001"],
        ["2026-09-14T13:37:51", 170000.00, 50, "", 169950.00, "EASYINFO fee for S11111111,2026-09-14", ""],
        ["2026-09-14T13:39:40", 169950.00, 50, "", 169900.00, "EASYINFO fee for BOM20260914133933528613,2026-09-1", ""],
        ["2026-09-15T10:38", 169900.00, 8500, "", 161400.00, "400000******1111:14-09-2026 13:39:34:BOM:SEZIS BUF", ""],
        ["2026-09-15T20:53:40", 161400.00, 5000, "", 156400.00, "SY5000 99112233!F", "8115006437"],
        ["2026-09-16T10:48:43", 156400.00, 5000, "", 151400.00, "400000******2222:15-09-2026 06:42:27:POS:UBCARD AP", ""],
        ["2026-09-21T10:04:51", 151400.00, "", 95000, 246400.00, "O-БАТ ДОРЖ", "5402600001"],
        ["2026-09-26T13:40:31", 246400.00, 97070, "", 149330.00, "400000******1111:25-09-2026 16:57:20:POS:MILD COSM", ""],
        ["2026-09-29T12:40:54", 149330.00, 200, "", 149130.00, "Charges for PORD Customer Payment :000658547340", "MN440032005006000001"],
        ["2026-09-29T12:40:54", 149130.00, 17500, "", 131630.00, "SocialPay гүйлгээ,ЦЭЦЭГ БОЛД,ХАС", "MN440032005006000001"],
        ["2026-09-29T13:27:25", 131630.00, "", 52500, 184130.00, "ШИЛЖҮҮЛЭГЧ ЦЭЦЭГ-ЦЭЦЭГ БОЛД", "5006000001"],
        ["2026-10-01T17:12:04", 184130.00, 72981.95, "", 111148.05, "400000******1111:29-09-2026 13:28:49:MCI:APPLE COM", ""],
        ["2026-10-04T13:08:04", 111108.05, 3960, "", 107148.05, "qpay 101782630891166, ccp-2494a138- 40a0", "MN680005005021900000"],
        ["2026-10-04T13:08:04", 111148.05, 40, "", 111108.05, "QPAY 335818730921540 ТЕСТ ХХК 4000 TRX", "2015113727"],
        ["2026-10-08T21:01:02", 107148.05, "", 110000, 217148.05, "O-БААТАР САРУУЛ", "5032300001"],
        ["2026-10-09T10:02:06", 217148.05, 106500, "", 110648.05, "400000******1111:08-10-2026 21:01:13:BOM:KFC ZURAG", ""],
        ["2026-10-09T12:57:13", 110648.05, 25350, "", 85298.05, "400000******1111:07-10-2026 14:12:55:POS:CU-281 KH", ""],
        ["2026-10-10T11:30:02", 85298.05, 3000, "", 82298.05, "99000000!OD3DAY!F", "2020001511"],
        ["2026-10-11T13:50:52", 82298.05, 4950, "", 77348.05, "400000******1111:09-10-2026 15:06:07:POS:CU-24 SHO", ""],
        ["2026-10-14T09:40:25", 77348.05, 12700, "", 64648.05, "400000******1111:13-10-2026 13:09:39:POS:SEZS - TS", ""],
    ]
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def db_session() -> Iterator[Session]:
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    session = factory()
    try:
        yield session
    finally:
        session.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_parser_extracts_statement_details_and_reconciles() -> None:
    parsed_workbook = load_workbook(BytesIO(_statement_bytes()), data_only=True)

    assert GolomtParser.detect(parsed_workbook) is True
    statement = detect_parser(parsed_workbook).parse(parsed_workbook)
    assert statement.account_no == "1800000001"
    assert statement.currency == "MNT"
    assert statement.iban == "MN000015001800000001"
    assert statement.holder_name == "ТЕСТ"
    assert statement.period_from.isoformat() == "2026-01-01"
    assert statement.period_to.isoformat() == "2026-01-14"
    assert len(statement.transactions) == 3
    assert statement.transactions[0].amount == Decimal("10.00")
    assert statement.transactions[0].direction == "out"
    assert statement.warnings == []


def test_parser_reports_balance_mismatch() -> None:
    workbook = load_workbook(BytesIO(_statement_bytes()), data_only=True)
    workbook.active["E3"] = 91.00
    workbook.active["E5"] = 86.00

    statement = GolomtParser().parse(workbook)

    assert any("Row 3: balance does not reconcile" in warning for warning in statement.warnings)
    assert any("Statement balance does not reconcile" in warning for warning in statement.warnings)


def test_full_sample_enrichment_and_fee_linking() -> None:
    workbook = load_workbook(BytesIO(_full_sample_bytes()), data_only=True)
    statement = GolomtParser().parse(workbook)
    items = [enrich(txn, {"5000000001"}) for txn in statement.transactions]
    link_fees(items)

    assert len(items) == 20
    assert statement.warnings == []
    assert sum((item.raw.amount for item in items if item.raw.direction == "in"), Decimal("0")) == Decimal("277500.00")
    assert sum((item.raw.amount for item in items if item.raw.direction == "out"), Decimal("0")) == Decimal("362851.95")
    assert Counter(item.txn_type for item in items) == {
        "card": 8,
        "fee": 4,
        "p2p_in": 3,
        "own_transfer": 1,
        "p2p_out": 1,
        "qpay": 1,
        "topup": 2,
    }
    assert normalize_account("MN440032005006000001") == "5006000001"
    assert items[3].merchant == "SEZIS BUF"
    assert items[3].txn_at == datetime(2026, 9, 14, 13, 39, 34)
    assert items[9].counterparty_key == items[10].counterparty_key == "A:5006000001"
    assert items[9].counterparty_name == "ЦЭЦЭГ БОЛД"
    assert items[1].parent_index == 0
    assert items[2].parent_index == 3
    assert items[8].parent_index == 9
    assert items[13].parent_index == 12


def test_import_is_idempotent_links_fee_and_classifies(db_session: Session) -> None:
    fee_category = Category(name="Bank fees", kind="expense", sort_order=10)
    db_session.add(fee_category)
    db_session.flush()
    db_session.add(
        Rule(
            priority=10,
            field="txn_type",
            match_type="equals",
            pattern="fee",
            category_id=fee_category.id,
            active=True,
        )
    )
    db_session.commit()
    data = _statement_bytes()

    first = import_file(db_session, data, "statement.xlsx")

    assert first.new_rows == 3
    assert first.duplicate_rows == 0
    assert first.warnings == []
    transactions = db_session.execute(
        select(Transaction).order_by(Transaction.posted_at, Transaction.id)
    ).scalars().all()
    socialpay = next(txn for txn in transactions if txn.description.startswith("SocialPay"))
    fee = next(txn for txn in transactions if txn.txn_type == "fee")
    card = next(txn for txn in transactions if txn.txn_type == "card")
    assert fee.parent_txn_id == socialpay.id
    assert fee.category_id == fee_category.id
    assert fee.category_source == "rule"
    assert card.txn_at == datetime(2026, 1, 2, 10, 2)
    assert card.posted_at == datetime(2026, 1, 2, 10, 3)

    repeated = import_file(db_session, data, "statement.xlsx")
    assert repeated.already_imported is True
    assert repeated.batch_id == first.batch_id
    assert repeated.bank_code == "GOLOMT"

    same_rows_different_file = import_file(
        db_session,
        _statement_bytes(report_date="2026-01-15"),
        "statement-copy.xlsx",
    )
    assert same_rows_different_file.new_rows == 0
    assert same_rows_different_file.duplicate_rows == 3
    assert db_session.execute(select(Transaction)).scalars().all()


def test_import_failure_rolls_back_without_creating_records(db_session: Session) -> None:
    workbook = Workbook()
    workbook.active.append(["not", "a", "Golomt", "statement"])
    buffer = BytesIO()
    workbook.save(buffer)

    with pytest.raises(ValueError, match="Unsupported statement format"):
        import_file(db_session, buffer.getvalue(), "unsupported.xlsx")

    assert db_session.execute(select(Transaction)).scalars().all() == []
